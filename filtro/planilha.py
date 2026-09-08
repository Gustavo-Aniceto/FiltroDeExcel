"""Leitura e escrita de planilhas (.xlsx, .xlsm, .csv, .tsv).

O módulo trabalha com uma estrutura simples (:class:`Planilha`): uma lista de
nomes de colunas e uma lista de linhas. Isso mantém o resto do sistema
independente do formato do arquivo de origem.
"""

from __future__ import annotations

import csv
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence

try:  # pragma: no cover - openpyxl só é necessário para arquivos Excel
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = load_workbook = None  # type: ignore[assignment]

EXTENSOES_EXCEL = {".xlsx", ".xlsm"}
EXTENSOES_TEXTO = {".csv", ".tsv", ".txt"}
FORMATOS_SUPORTADOS = sorted(EXTENSOES_EXCEL | EXTENSOES_TEXTO)

LARGURA_MAXIMA_COLUNA = 50


class ErroDePlanilha(Exception):
    """Erro ao ler, interpretar ou gravar uma planilha."""


def chave(nome: Any) -> str:
    """Normaliza um nome de coluna: sem acentos, minúsculo e sem espaços extras."""
    texto = "" if nome is None else str(nome)
    decomposto = unicodedata.normalize("NFKD", texto)
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return " ".join(sem_acento.lower().split())


@dataclass
class Planilha:
    """Uma tabela em memória: cabeçalho + linhas."""

    colunas: list[str]
    linhas: list[list[Any]] = field(default_factory=list)
    nome: str = "Planilha1"
    delimitador: str = ","

    def __len__(self) -> int:
        return len(self.linhas)

    def indice(self, coluna: str) -> int:
        """Devolve a posição de ``coluna``, ignorando acentos e maiúsculas."""
        if coluna in self.colunas:
            return self.colunas.index(coluna)
        procurada = chave(coluna)
        for posicao, nome in enumerate(self.colunas):
            if chave(nome) == procurada:
                return posicao
        disponiveis = ", ".join(self.colunas) or "(nenhuma)"
        raise ErroDePlanilha(
            f"coluna {coluna!r} não existe na planilha. Colunas disponíveis: {disponiveis}"
        )

    def valor(self, linha: Sequence[Any], coluna: str) -> Any:
        posicao = self.indice(coluna)
        return linha[posicao] if posicao < len(linha) else None

    def selecionar(self, colunas: Sequence[str]) -> "Planilha":
        """Cria uma nova planilha apenas com as colunas indicadas, na ordem dada."""
        posicoes = [self.indice(c) for c in colunas]
        nomes = [self.colunas[p] for p in posicoes]
        linhas = [[linha[p] if p < len(linha) else None for p in posicoes] for linha in self.linhas]
        return Planilha(nomes, linhas, self.nome, self.delimitador)

    def com_linhas(self, linhas: list[list[Any]]) -> "Planilha":
        return Planilha(list(self.colunas), linhas, self.nome, self.delimitador)


# ---------------------------------------------------------------- leitura


def ler(caminho: str | Path, aba: str | int | None = None, linha_cabecalho: int = 1) -> Planilha:
    """Lê uma planilha de disco e devolve uma :class:`Planilha`.

    ``aba`` só se aplica a arquivos Excel (nome ou índice começando em 1).
    ``linha_cabecalho`` indica em qual linha (1 = primeira) estão os títulos.
    """
    caminho = Path(caminho)
    if not caminho.exists():
        raise ErroDePlanilha(f"arquivo não encontrado: {caminho}")
    if linha_cabecalho < 1:
        raise ErroDePlanilha("a linha do cabeçalho precisa ser 1 ou maior")

    sufixo = caminho.suffix.lower()
    if sufixo in EXTENSOES_EXCEL:
        return _ler_excel(caminho, aba, linha_cabecalho)
    if sufixo in EXTENSOES_TEXTO:
        return _ler_texto(caminho, linha_cabecalho)
    raise ErroDePlanilha(
        f"formato não suportado: {sufixo or caminho.name!r}. "
        f"Use um destes: {', '.join(FORMATOS_SUPORTADOS)}"
    )


def _ler_excel(caminho: Path, aba: str | int | None, linha_cabecalho: int) -> Planilha:
    if load_workbook is None:  # pragma: no cover
        raise ErroDePlanilha("openpyxl não está instalado; rode: pip install openpyxl")
    try:
        arquivo = load_workbook(caminho, data_only=True, read_only=True)
    except Exception as erro:  # pragma: no cover - depende do arquivo
        raise ErroDePlanilha(f"não foi possível abrir {caminho}: {erro}") from erro

    try:
        planilha_excel = _escolher_aba(arquivo, aba)
        linhas_brutas = [list(linha) for linha in planilha_excel.iter_rows(values_only=True)]
        nome_aba = planilha_excel.title
    finally:
        arquivo.close()

    colunas, linhas = _separar_cabecalho(linhas_brutas, linha_cabecalho, caminho)
    return Planilha(colunas, linhas, nome_aba)


def _escolher_aba(arquivo, aba):
    if aba is None:
        return arquivo.active
    if isinstance(aba, int) or (isinstance(aba, str) and aba.isdigit()):
        numero = int(aba)
        if not 1 <= numero <= len(arquivo.sheetnames):
            raise ErroDePlanilha(
                f"aba {numero} não existe. Abas: {', '.join(arquivo.sheetnames)}"
            )
        return arquivo[arquivo.sheetnames[numero - 1]]
    procurada = chave(aba)
    for nome in arquivo.sheetnames:
        if chave(nome) == procurada:
            return arquivo[nome]
    raise ErroDePlanilha(f"aba {aba!r} não existe. Abas: {', '.join(arquivo.sheetnames)}")


def _ler_texto(caminho: Path, linha_cabecalho: int) -> Planilha:
    conteudo = _ler_com_codificacao(caminho)
    delimitador = _detectar_delimitador(conteudo, caminho.suffix.lower())
    linhas_brutas = [list(linha) for linha in csv.reader(conteudo.splitlines(), delimiter=delimitador)]
    colunas, linhas = _separar_cabecalho(linhas_brutas, linha_cabecalho, caminho)
    return Planilha(colunas, linhas, caminho.stem, delimitador)


def _ler_com_codificacao(caminho: Path) -> str:
    for codificacao in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return caminho.read_text(encoding=codificacao)
        except UnicodeDecodeError:
            continue
    raise ErroDePlanilha(f"não foi possível decodificar o arquivo {caminho}")


def _detectar_delimitador(conteudo: str, sufixo: str) -> str:
    if sufixo == ".tsv":
        return "\t"
    amostra = "\n".join(conteudo.splitlines()[:20])
    if not amostra.strip():
        return ","
    contagens = {sep: amostra.count(sep) for sep in (";", ",", "\t", "|")}
    melhor = max(contagens, key=lambda sep: contagens[sep])
    return melhor if contagens[melhor] else ","


def _separar_cabecalho(
    linhas_brutas: list[list[Any]], linha_cabecalho: int, caminho: Path
) -> tuple[list[str], list[list[Any]]]:
    if len(linhas_brutas) < linha_cabecalho:
        raise ErroDePlanilha(
            f"{caminho} não tem a linha {linha_cabecalho} para usar como cabeçalho"
        )
    cabecalho = linhas_brutas[linha_cabecalho - 1]
    while cabecalho and _vazio(cabecalho[-1]):
        cabecalho.pop()
    if not cabecalho:
        raise ErroDePlanilha(f"o cabeçalho de {caminho} está vazio")

    colunas = _nomear_colunas(cabecalho)
    largura = len(colunas)
    linhas = []
    for linha in linhas_brutas[linha_cabecalho:]:
        if all(_vazio(celula) for celula in linha):
            continue
        linha = list(linha[:largura])
        linha.extend([None] * (largura - len(linha)))
        linhas.append(linha)
    return colunas, linhas


def _nomear_colunas(cabecalho: Sequence[Any]) -> list[str]:
    nomes: list[str] = []
    usados: dict[str, int] = {}
    for posicao, bruto in enumerate(cabecalho, start=1):
        nome = "" if bruto is None else str(bruto).strip()
        if not nome:
            nome = f"Coluna {posicao}"
        base = chave(nome)
        if base in usados:
            usados[base] += 1
            nome = f"{nome} ({usados[base]})"
        else:
            usados[base] = 1
        nomes.append(nome)
    return nomes


def _vazio(celula: Any) -> bool:
    return celula is None or (isinstance(celula, str) and not celula.strip())


# ---------------------------------------------------------------- escrita


def escrever(planilha: Planilha, caminho: str | Path, delimitador: str | None = None) -> Path:
    """Grava a planilha. O formato é definido pela extensão de ``caminho``."""
    caminho = Path(caminho)
    sufixo = caminho.suffix.lower()
    if caminho.parent and not caminho.parent.exists():
        caminho.parent.mkdir(parents=True, exist_ok=True)

    if sufixo in EXTENSOES_EXCEL:
        _escrever_excel(planilha, caminho)
    elif sufixo in EXTENSOES_TEXTO:
        _escrever_texto(planilha, caminho, delimitador or planilha.delimitador)
    else:
        raise ErroDePlanilha(
            f"formato de saída não suportado: {sufixo or caminho.name!r}. "
            f"Use um destes: {', '.join(FORMATOS_SUPORTADOS)}"
        )
    return caminho


def _escrever_excel(planilha: Planilha, caminho: Path) -> None:
    if Workbook is None:  # pragma: no cover
        raise ErroDePlanilha("openpyxl não está instalado; rode: pip install openpyxl")
    arquivo = Workbook()
    aba = arquivo.active
    aba.title = (planilha.nome or "Planilha1")[:31]
    aba.append(planilha.colunas)
    for linha in planilha.linhas:
        aba.append([_valor_para_excel(celula) for celula in linha])

    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"
    if planilha.colunas:
        aba.auto_filter.ref = f"A1:{get_column_letter(len(planilha.colunas))}{len(planilha.linhas) + 1}"
    for posicao, nome in enumerate(planilha.colunas, start=1):
        conteudos = [len(str(nome))]
        conteudos += [len(_texto(linha[posicao - 1])) for linha in planilha.linhas[:200]]
        largura = min(max(conteudos) + 2, LARGURA_MAXIMA_COLUNA)
        aba.column_dimensions[get_column_letter(posicao)].width = largura
    arquivo.save(caminho)


def _escrever_texto(planilha: Planilha, caminho: Path, delimitador: str) -> None:
    with caminho.open("w", encoding="utf-8-sig", newline="") as saida:
        escritor = csv.writer(saida, delimiter=delimitador)
        escritor.writerow(planilha.colunas)
        for linha in planilha.linhas:
            escritor.writerow([_texto(celula) for celula in linha])


def _valor_para_excel(celula: Any) -> Any:
    if isinstance(celula, (int, float, datetime, date, bool)) or celula is None:
        return celula
    return str(celula)


def _texto(celula: Any) -> str:
    """Converte uma célula em texto de forma previsível (datas em ISO, inteiros sem .0)."""
    if celula is None:
        return ""
    if isinstance(celula, bool):
        return "sim" if celula else "não"
    if isinstance(celula, float) and celula.is_integer():
        return str(int(celula))
    if isinstance(celula, datetime):
        if (celula.hour, celula.minute, celula.second) == (0, 0, 0):
            return celula.date().isoformat()
        return celula.isoformat(sep=" ")
    if isinstance(celula, date):
        return celula.isoformat()
    return str(celula)
